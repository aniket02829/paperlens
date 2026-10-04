"""Build data/paper_classifier.db from the public ranking and indexing lists.

Usage:
    python setup_db.py                                   # build everything from data/sources/
    python setup_db.py --conferences-only <core_csv>     # rebuild only the CORE table
    python update_data.py                                # download fresh lists, then build

Sources (see data/sources/README.md for provenance and licences):
    sjr.csv              Scimago Journal Rank export (';'-separated, comma decimals)
    core.csv             CORE conference ranking export; it has NO header row:
                         id, title, acronym, source, rank, primary, FoR1, FoR2, FoR3
    doaj.csv             Directory of Open Access Journals, full CSV export
    scopus_sources.xlsx  Scopus source title list (sheets 'Scopus Sources ...',
                         'Discontinued Titles ...', 'Serial Conf. Proc. with Profile')
    rw_hijacked.csv      Retraction Watch Hijacked Journal Checker (Google Sheet export)
    spj_journals.csv     Stop Predatory Journals: journals
    spj_publishers.csv   Stop Predatory Journals: publishers

Every file is optional; a missing one leaves its table empty and prints a warning.
Journal ids are Scimago source ids, so they stay stable across editions.
"""
import csv
import glob
import os
import re
import sqlite3
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from classifier.db_base import name_key  # noqa: E402

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'data', 'paper_classifier.db')
SOURCES_DIR = os.path.join(BASE_DIR, 'data', 'sources')

csv.field_size_limit(10 ** 8)


def normalize_issn(issn: str) -> str:
    return re.sub(r'[^0-9X]', '', (issn or '').upper())


def split_issns(text: str):
    """'1941-6520,  1941-6067' -> ['19416520', '19416067']"""
    out = []
    for part in re.split(r'[,;/\s]+', text or ''):
        issn = normalize_issn(part)
        if len(issn) == 8 and issn not in out:
            out.append(issn)
    return out


def find_source(*names: str) -> str:
    """The first existing file in data/sources/ matching one of the names or globs."""
    for name in names:
        matches = sorted(glob.glob(os.path.join(SOURCES_DIR, name)), key=os.path.getmtime, reverse=True)
        if matches:
            return matches[0]
    return ''


# ----------------------------------------------------------------------------- schema

SCHEMA = '''
CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT);

CREATE TABLE journals (
    id INTEGER PRIMARY KEY,            -- Scimago source id
    title TEXT, title_normalized TEXT, issn_list TEXT,
    source_type TEXT,                  -- journal | book series | conference and proceedings | trade journal
    sjr_score REAL, best_quartile TEXT, h_index INTEGER,
    cites_per_doc REAL, total_docs INTEGER, coverage TEXT,
    publisher TEXT, country TEXT, categories TEXT, areas TEXT, open_access TEXT
);
CREATE TABLE journal_issns (journal_id INTEGER, issn TEXT);
CREATE INDEX idx_journal_issns_issn ON journal_issns(issn);
CREATE INDEX idx_journals_title ON journals(title_normalized);

CREATE TABLE doaj_journals (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT, title_normalized TEXT, alt_title TEXT, url TEXT,
    publisher TEXT, country TEXT,
    apc TEXT, apc_amount TEXT, review_process TEXT, weeks_to_publication TEXT,
    license TEXT, subjects TEXT, added_on TEXT, article_count INTEGER
);
CREATE TABLE doaj_issns (journal_id INTEGER, issn TEXT);
CREATE INDEX idx_doaj_issns_issn ON doaj_issns(issn);
CREATE INDEX idx_doaj_title ON doaj_journals(title_normalized);

CREATE TABLE scopus_sources (
    source_id INTEGER PRIMARY KEY,
    title TEXT, title_normalized TEXT, issn TEXT, eissn TEXT,
    source_type TEXT, active INTEGER, discontinued INTEGER,
    discontinued_year TEXT, discontinued_reason TEXT,
    coverage TEXT, open_access TEXT, publisher TEXT, title_history TEXT
);
CREATE INDEX idx_scopus_issn ON scopus_sources(issn);
CREATE INDEX idx_scopus_eissn ON scopus_sources(eissn);
CREATE INDEX idx_scopus_title ON scopus_sources(title_normalized);

-- Journals and publishers flagged by public lists: kind is 'hijacked',
-- 'predatory_journal' or 'predatory_publisher'. For hijacked entries, name/url/issns
-- describe the clone and authentic_* the real journal.
CREATE TABLE watchlist (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    kind TEXT, name TEXT, name_normalized TEXT, url TEXT, issns TEXT,
    authentic_name TEXT, authentic_url TEXT, authentic_issns TEXT, source TEXT
);
CREATE TABLE watchlist_issns (watchlist_id INTEGER, issn TEXT);
CREATE INDEX idx_watchlist_issns_issn ON watchlist_issns(issn);
CREATE INDEX idx_watchlist_name ON watchlist(name_normalized);
'''

CONFERENCE_SCHEMA = '''
CREATE TABLE conferences (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT, title_normalized TEXT, acronym TEXT, acronym_normalized TEXT,
    source TEXT, rank TEXT, is_primary TEXT, for_code_1 TEXT, for_code_2 TEXT, for_code_3 TEXT
);
CREATE INDEX idx_conferences_acronym ON conferences(acronym_normalized);
CREATE INDEX idx_conferences_title ON conferences(title_normalized);
'''


def set_meta(conn, key, value):
    conn.execute('INSERT OR REPLACE INTO meta (key, value) VALUES (?, ?)', (key, str(value)))


# ----------------------------------------------------------------------------- Scimago

def load_sjr(conn, path) -> int:
    if not path or not os.path.exists(path):
        print(f"[!] SJR CSV not found ({path}); journals table left empty")
        return 0
    count = 0
    with open(path, encoding='utf-8-sig', errors='replace', newline='') as f:
        reader = csv.DictReader(f, delimiter=';')
        docs_column = next((c for c in reader.fieldnames or [] if c.startswith('Total Docs. (')), None)
        edition = re.search(r'\((\d{4})\)', docs_column or '')
        if edition:
            set_meta(conn, 'sjr_edition', edition.group(1))

        def number(text):
            try:
                return float((text or '').strip().replace(',', '.'))
            except ValueError:
                return None

        for row in reader:
            title = (row.get('Title') or '').strip()
            source_id = (row.get('Sourceid') or '').strip()
            if not title or not source_id.isdigit():
                continue
            issns = split_issns(row.get('Issn', ''))
            conn.execute('''INSERT OR REPLACE INTO journals (id, title, title_normalized, issn_list, source_type,
                sjr_score, best_quartile, h_index, cites_per_doc, total_docs, coverage,
                publisher, country, categories, areas, open_access) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
                int(source_id), title, title.lower(), ', '.join(issns), (row.get('Type') or '').strip().lower(),
                number(row.get('SJR')) or 0.0, (row.get('SJR Best Quartile') or '').strip(),
                int(number(row.get('H index')) or 0), number(row.get('Citations / Doc. (2years)')),
                int(number(row.get(docs_column) if docs_column else None) or 0), (row.get('Coverage') or '').strip(),
                (row.get('Publisher') or '').strip(), (row.get('Country') or '').strip(),
                (row.get('Categories') or '').strip(), (row.get('Areas') or '').strip(),
                (row.get('Open Access') or '').strip()))
            conn.executemany('INSERT INTO journal_issns (journal_id, issn) VALUES (?, ?)',
                             [(int(source_id), issn) for issn in issns])
            count += 1
    return count


# ----------------------------------------------------------------------------- CORE

CORE_COLUMNS = ['id', 'Title', 'Acronym', 'Source', 'Rank', 'Primary', 'FoR1', 'FoR2', 'FoR3']


def read_core_rows(f):
    """Yield CORE rows as dicts, whether or not the CSV has a header row."""
    reader = csv.reader(f, delimiter=',')
    first = next(reader, None)
    if first is None:
        return
    if 'Title' in [c.strip() for c in first]:
        header = [c.strip() for c in first]
    else:
        header = CORE_COLUMNS
        yield dict(zip(header, first))
    for row in reader:
        yield dict(zip(header, row))


def load_core(conn, path) -> int:
    if not path or not os.path.exists(path):
        print(f"[!] CORE CSV not found ({path}); conferences table left empty")
        return 0
    count = 0
    with open(path, encoding='utf-8-sig', errors='replace', newline='') as f:
        for row in read_core_rows(f):
            title = (row.get('Title') or '').strip()
            acronym = (row.get('Acronym') or '').strip()
            if not title and not acronym:
                continue
            conn.execute('''INSERT INTO conferences (title, title_normalized, acronym, acronym_normalized,
                source, rank, is_primary, for_code_1, for_code_2, for_code_3) VALUES (?,?,?,?,?,?,?,?,?,?)''', (
                title, title.lower(), acronym, acronym.lower(), (row.get('Source') or '').strip(),
                (row.get('Rank') or '').strip(), (row.get('Primary') or '').strip(),
                (row.get('FoR1') or '').strip(), (row.get('FoR2') or '').strip(), (row.get('FoR3') or '').strip()))
            count += 1
    source = conn.execute("SELECT source FROM conferences WHERE source != '' LIMIT 1").fetchone()
    if source:
        set_meta(conn, 'core_edition', source[0])
    elif re.search(r'(\d{4})', os.path.basename(path)):
        set_meta(conn, 'core_edition', re.search(r'(\d{4})', os.path.basename(path)).group(1))
    return count


# ----------------------------------------------------------------------------- DOAJ

def load_doaj(conn, path) -> int:
    if not path or not os.path.exists(path):
        print(f"[!] DOAJ CSV not found ({path}); DOAJ table left empty")
        return 0
    count = 0
    with open(path, encoding='utf-8-sig', errors='replace', newline='') as f:
        for row in csv.DictReader(f):
            title = (row.get('Journal title') or '').strip()
            if not title:
                continue
            issns = split_issns(f"{row.get('Journal ISSN (print version)', '')} {row.get('Journal EISSN (online version)', '')}")
            cur = conn.execute('''INSERT INTO doaj_journals (title, title_normalized, alt_title, url, publisher, country,
                apc, apc_amount, review_process, weeks_to_publication, license, subjects, added_on, article_count)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
                title, name_key(title), (row.get('Alternative title') or '').strip(), (row.get('Journal URL') or '').strip(),
                (row.get('Publisher') or '').strip(), (row.get('Country of publisher') or '').strip(),
                (row.get('APC') or '').strip(), (row.get('APC amount') or '').strip(),
                (row.get('Review process') or '').strip(),
                (row.get('Average number of weeks between article submission and publication') or '').strip(),
                (row.get('Journal license') or '').strip(), (row.get('Subjects') or '').strip(),
                (row.get('Added on Date') or '')[:10], int(row.get('Number of Article Records') or 0 or 0)))
            conn.executemany('INSERT INTO doaj_issns (journal_id, issn) VALUES (?, ?)', [(cur.lastrowid, i) for i in issns])
            count += 1
    set_meta(conn, 'doaj_downloaded', time.strftime('%Y-%m-%d', time.localtime(os.path.getmtime(path))))
    return count


# ----------------------------------------------------------------------------- Scopus

def load_scopus(conn, path) -> int:
    if not path or not os.path.exists(path):
        print(f"[!] Scopus source list not found ({path}); Scopus table left empty")
        return 0
    try:
        import openpyxl
    except ImportError:
        print("[!] openpyxl is needed to read the Scopus list: pip install -r requirements-dev.txt")
        return 0
    wb = openpyxl.load_workbook(path, read_only=True)

    def sheet(prefix):
        return next((wb[n] for n in wb.sheetnames if n.lower().startswith(prefix.lower())), None)

    sources = sheet('Scopus Sources')
    if sources is None:
        print("[!] The workbook has no 'Scopus Sources' sheet")
        return 0
    edition = sources.title.replace('Scopus Sources', '').strip().rstrip('.')
    set_meta(conn, 'scopus_list', edition.replace('Aug.', 'August').replace('Jun.', 'June').replace('Oct.', 'October')
             .replace('Feb.', 'February').replace('Apr.', 'April').replace('Dec.', 'December'))

    # Why a title left Scopus, from the 'Discontinued Titles' sheet (two header rows).
    discontinued = {}
    disc_sheet = sheet('Discontinued Titles')
    if disc_sheet is not None:
        rows = disc_sheet.iter_rows(values_only=True)
        header = None
        for row in rows:
            if header is None:
                if row and row[0] == 'Sourcerecord ID':
                    header = [str(c or '').strip() for c in row]
                continue
            item = dict(zip(header, row))
            source_id = str(item.get('Sourcerecord ID') or '').strip()
            if source_id.isdigit():
                year = str(item.get('Year') or '').strip()
                discontinued[int(source_id)] = (year if re.fullmatch(r'\d{4}', year) else '',
                                                str(item.get('Indexation Change') or '').strip())

    count = 0
    rows = sources.iter_rows(values_only=True)
    header = [str(c or '').strip() for c in next(rows)]
    for row in rows:
        item = dict(zip(header, row))
        source_id = str(item.get('Sourcerecord ID') or '').strip()
        title = str(item.get('Source Title') or '').strip()
        if not source_id.isdigit() or not title:
            continue
        year, reason = discontinued.get(int(source_id), ('', ''))
        flagged = bool(str(item.get('Titles Discontinued by Scopus') or '').strip()) or int(source_id) in discontinued
        conn.execute('''INSERT OR REPLACE INTO scopus_sources (source_id, title, title_normalized, issn, eissn, source_type,
            active, discontinued, discontinued_year, discontinued_reason, coverage, open_access, publisher, title_history)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
            int(source_id), title, name_key(title), normalize_issn(str(item.get('ISSN') or '')),
            normalize_issn(str(item.get('EISSN') or '')), str(item.get('Source Type') or '').strip(),
            1 if str(item.get('Active or Inactive') or '').strip().lower() == 'active' else 0,
            1 if flagged else 0, year, reason, str(item.get('Coverage') or '').strip(),
            str(item.get('Open Access Status') or '').strip(), str(item.get('Publisher') or '').strip(),
            str(item.get('Title History Indication') or '').strip()))
        count += 1

    proceedings = sheet('Serial Conf. Proc')
    if proceedings is not None:
        rows = proceedings.iter_rows(values_only=True)
        header = [str(c or '').strip() for c in next(rows)]
        for row in rows:
            item = dict(zip(header, row))
            source_id = str(item.get('Sourcerecord ID') or '').strip()
            title = str(item.get('Source Title') or '').strip()
            if not source_id.isdigit() or not title:
                continue
            conn.execute('''INSERT OR IGNORE INTO scopus_sources (source_id, title, title_normalized, issn, eissn, source_type,
                active, discontinued, discontinued_year, discontinued_reason, coverage, open_access, publisher, title_history)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
                int(source_id), title, name_key(title), normalize_issn(str(item.get('ISSN') or '')),
                normalize_issn(str(item.get('EISSN') or '')), 'Conference Proceedings', 1,
                1 if str(item.get('Titles Discontinued by Scopus') or '').strip() else 0, '', '',
                str(item.get('Coverage') or '').strip(), '', '', ''))
            count += 1
    return count


# ----------------------------------------------------------------------------- watchlists

def load_hijacked(conn, path) -> int:
    """Retraction Watch Hijacked Journal Checker: the sheet has two preamble rows."""
    if not path or not os.path.exists(path):
        print(f"[!] Hijacked journal list not found ({path})")
        return 0
    count = 0
    with open(path, encoding='utf-8-sig', errors='replace', newline='') as f:
        header = None
        for row in csv.reader(f):
            if header is None:
                if row and row[0].startswith('First created'):
                    updated = re.search(r'last updated\s+(.+)$', row[0])
                    if updated:
                        set_meta(conn, 'hijacked_updated', updated.group(1).strip())
                if 'Hijacked Journal Title' in [c.strip() for c in row]:
                    header = [c.strip() for c in row]
                continue
            item = dict(zip(header, row))
            name = (item.get('Hijacked Journal Title') or '').strip()
            authentic = (item.get('Original journal') or '').strip()
            if not name and not authentic:
                continue
            clone_issns = split_issns(item.get('ISSN (Hijacked)', ''))
            real_issns = split_issns(item.get('ISSN (Original)', ''))
            cur = conn.execute('''INSERT INTO watchlist (kind, name, name_normalized, url, issns, authentic_name,
                authentic_url, authentic_issns, source) VALUES ('hijacked',?,?,?,?,?,?,?,?)''', (
                name, name_key(name), (item.get('URL (Hijacked)') or '').strip(), ', '.join(clone_issns),
                authentic, (item.get('URL (Original Journal)') or '').strip(), ', '.join(real_issns),
                'Retraction Watch Hijacked Journal Checker'))
            conn.executemany('INSERT INTO watchlist_issns (watchlist_id, issn) VALUES (?, ?)',
                             [(cur.lastrowid, i) for i in dict.fromkeys(clone_issns + real_issns)])
            count += 1
    return count


def load_predatory(conn, path, kind) -> int:
    """Stop Predatory Journals lists (url,name,abbr)."""
    if not path or not os.path.exists(path):
        print(f"[!] Predatory list not found ({path})")
        return 0
    count = 0
    with open(path, encoding='utf-8-sig', errors='replace', newline='') as f:
        for row in csv.DictReader(f):
            name = (row.get('name') or '').strip()
            if not name:
                continue
            conn.execute('''INSERT INTO watchlist (kind, name, name_normalized, url, issns, authentic_name,
                authentic_url, authentic_issns, source) VALUES (?,?,?,?,'','','','',?)''', (
                kind, name, name_key(name), (row.get('url') or '').strip(), 'Stop Predatory Journals'))
            count += 1
    return count


# ----------------------------------------------------------------------------- building

def build(sjr=None, core=None, doaj=None, scopus=None, hijacked=None, spj_journals=None, spj_publishers=None):
    sjr = sjr or find_source('sjr*.csv', 'scimagojr*.csv')
    core = core or find_source('CORE*.csv', 'core*.csv')
    doaj = doaj or find_source('doaj*.csv')
    scopus = scopus or find_source('scopus*.xlsx', 'ext_list*.xlsx')
    hijacked = hijacked or find_source('rw_hijacked*.csv')
    spj_journals = spj_journals or find_source('spj_journals.csv')
    spj_publishers = spj_publishers or find_source('spj_publishers.csv')

    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    tmp_path = DB_PATH + '.building'
    if os.path.exists(tmp_path):
        os.remove(tmp_path)
    conn = sqlite3.connect(tmp_path)
    try:
        conn.executescript(SCHEMA + CONFERENCE_SCHEMA)
        counts = {
            'journals': load_sjr(conn, sjr),
            'conferences': load_core(conn, core),
            'doaj': load_doaj(conn, doaj),
            'scopus': load_scopus(conn, scopus),
            'hijacked': load_hijacked(conn, hijacked),
            'predatory_journals': load_predatory(conn, spj_journals, 'predatory_journal'),
            'predatory_publishers': load_predatory(conn, spj_publishers, 'predatory_publisher'),
        }
        set_meta(conn, 'built_on', time.strftime('%Y-%m-%d'))
        conn.commit()
        conn.execute('VACUUM')
    finally:
        conn.close()
    os.replace(tmp_path, DB_PATH)
    for name, n in counts.items():
        print(f"  {name:22} {n:>8,}")
    print(f"Database written to {DB_PATH}")
    return counts


def rebuild_conferences_only(core_path: str):
    """Replace only the conferences table, leaving every other table untouched."""
    conn = sqlite3.connect(DB_PATH)
    try:
        conn.execute('DROP TABLE IF EXISTS conferences')
        conn.executescript(CONFERENCE_SCHEMA)
        conn.execute('CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT)')
        count = load_core(conn, core_path)
        conn.commit()
        print(f"[OK] Loaded {count:,} conferences into {DB_PATH}")
    finally:
        conn.close()


def main():
    if len(sys.argv) == 3 and sys.argv[1] == '--conferences-only':
        rebuild_conferences_only(sys.argv[2])
        return
    if len(sys.argv) == 3 and not sys.argv[1].startswith('-'):
        build(sjr=sys.argv[1], core=sys.argv[2])
        return
    build()


if __name__ == '__main__':
    main()
